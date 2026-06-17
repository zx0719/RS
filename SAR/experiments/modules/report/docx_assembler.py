"""
docx_assembler.py — M6 Word 文档组装模块

class DocxAssembler:
    assemble(evidence_package, output_dir, template_path=None) -> str

文档结构（对标成品.docx模板格式）：
  [首页表头_中]  航天通报
  [首页表头_下]  YYYY年MM月DD日
  [Normal 16pt]  正文内容
  [Normal]        附件列表（附件1/2/3标题行）
  附件1：侦察图（若有）
  附件2：组成分布统计表  [序号, 目标类型, 经度, 纬度]
  附件3：装备分布统计表  [序号, 目标类型, 经度, 纬度]
  页脚：卫星/传感器/成像时间/pipeline_run_id/acceptance_run_id + 自动生成标识（审核版）

支持两种模式：
  - 审核版（draft_mode=True）：页脚含"【自动生成草稿，待人工审核】"
  - 正式版（draft_mode=False）：无自动生成标识
"""

from __future__ import annotations

import os
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

try:
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Cm, Pt, RGBColor  # type: ignore
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

# ---------------------------------------------------------------------------
# 默认模板路径（优先使用需求目录中的真实成品.docx）
# ---------------------------------------------------------------------------

_TEMPLATE_CANDIDATES = [
    "/home/zhuxiang/RS/SAR/需求/成品.docx",
    "/home/zhuxiang/RS/SAR/文档/成品.docx",
]

DEFAULT_TEMPLATE_PATH = next(
    (p for p in _TEMPLATE_CANDIDATES if Path(p).exists()), None
)

# ---------------------------------------------------------------------------
# 表格列头（4列，对标成品.docx）
# ---------------------------------------------------------------------------

_TABLE_HEADERS = ["序号", "目标类型", "经度", "纬度"]
_COORD_PLACEHOLDER = "\u2014"
_MAX_TABLE_ROWS_PER_CHUNK = 28
_MIN_FINAL_TABLE_ROWS = 1
_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_XML_NS = {"w": _WORD_NS}


def _uri_to_local_path(uri: str) -> str | None:
    if uri and uri.startswith("file://"):
        return uri[len("file://"):]
    return None


def _format_coordinate(value: Any, coord_kind: str) -> str:
    if value is None:
        return _COORD_PLACEHOLDER

    try:
        coord = float(value)
    except (TypeError, ValueError):
        return _COORD_PLACEHOLDER

    if coord_kind == "lon" and -180.0 <= coord <= 180.0:
        return f"{coord:.9f}"
    if coord_kind == "lat" and -90.0 <= coord <= 90.0:
        return f"{coord:.9f}"
    return _COORD_PLACEHOLDER


def _chunk_table_rows(rows: list[dict]) -> list[list[dict]]:
    chunks = [
        rows[i : i + _MAX_TABLE_ROWS_PER_CHUNK]
        for i in range(0, len(rows), _MAX_TABLE_ROWS_PER_CHUNK)
    ]
    if (
        len(chunks) > 1
        and len(chunks[-1]) < _MIN_FINAL_TABLE_ROWS
        and len(chunks[-2]) + len(chunks[-1]) <= _MAX_TABLE_ROWS_PER_CHUNK + 3
    ):
        chunks[-2].extend(chunks[-1])
        chunks.pop()
    return chunks


class DocxAssembler:
    """M6 Word 文档组装器，输出格式对标成品.docx模板。"""

    TEMPLATE_VERSION = "brief-template-v2"

    def assemble(
        self,
        evidence_package: dict,
        output_dir: str,
        template_path: str | None = None,
        draft_mode: bool = True,
        include_image: bool = True,
    ) -> str:
        if not _DOCX_AVAILABLE:
            raise ImportError("python-docx 未安装，请执行：pip install python-docx")

        os.makedirs(output_dir, exist_ok=True)

        if template_path is None and DEFAULT_TEMPLATE_PATH:
            template_path = DEFAULT_TEMPLATE_PATH

        # 提取字段
        pkg = evidence_package
        report = pkg.get("report", {})
        objects = pkg.get("objects", [])
        attachments = pkg.get("attachments", {})
        inp = pkg.get("input", {})
        metadata = inp.get("metadata", {})
        mission = inp.get("mission", {})
        trace = pkg.get("trace", {})

        region_name = mission.get("region_name", "目标区域")
        report_date = report.get("report_date", "")
        body = report.get("body")
        if not isinstance(body, str) or not body.strip():
            raise ValueError("report.body is required before DOCX assembly.")
        tables = report.get("tables", {})
        component_table = tables.get("component_table", [])
        equipment_table = tables.get("equipment_table", [])

        # 若 tables 为空，从 objects 重建
        if not component_table and not equipment_table and objects:
            from .table_builder import TableBuilder
            tb = TableBuilder()
            component_table = tb.build_component_table(objects)
            equipment_table = tb.build_equipment_table(objects)

        # 格式化日期
        display_date = self._format_date(report_date, metadata.get("acquisition_time", ""))

        # 构建输出文件名
        package_id = pkg.get("package_id", "report")
        safe_id = re.sub(r"[^\w\-]", "_", package_id)
        draft_suffix = "_draft" if draft_mode else "_final"
        filename = f"{safe_id}{draft_suffix}.docx"
        output_path = str(Path(output_dir) / filename)

        # 加载文档（从模板清空正文，或新建）
        if template_path and Path(template_path).exists():
            doc = Document(template_path)
            self._clear_body_content(doc)
        else:
            doc = Document()

        self._set_margins(doc)

        # ---------- 标题：航天通报 ----------
        self._add_header_paragraphs(doc, display_date)

        # ---------- 正文 ----------
        self._add_body_text(doc, body)

        # ---------- 附件目录行（附件1/2/3引用行） ----------
        self._add_attachment_index(doc, region_name, include_image=include_image)

        # ---------- 附件1：侦察图 ----------
        if include_image:
            annotated = attachments.get("annotated_image", {})
            ann_uri = annotated.get("uri", "")
            ann_path = _uri_to_local_path(ann_uri) if ann_uri else None
            self._add_attachment1_image(doc, ann_path, region_name)

        # ---------- 附件2：组成分布统计表 ----------
        if include_image:
            doc.add_page_break()
        self._add_attachment2_table(doc, component_table, region_name)

        # ---------- 附件3：装备分布统计表 ----------
        doc.add_page_break()
        self._add_attachment3_table(doc, equipment_table, region_name)

        self._add_footer(
            doc,
            satellite=str(metadata.get("satellite", "")),
            sensor=str(metadata.get("sensor", "")),
            acq_time=str(metadata.get("acquisition_time", "")),
            pipeline_run_id=str(trace.get("pipeline_run_id", "")),
            acceptance_run_id=str(trace.get("acceptance_run_id", "")),
            draft_mode=draft_mode,
        )

        doc.save(output_path)
        self._strip_unreferenced_document_images(output_path)
        return output_path

    # ------------------------------------------------------------------
    # 内部组装方法
    # ------------------------------------------------------------------

    @staticmethod
    def _clear_body_content(doc: Any) -> None:
        """清除正文所有段落和表格，保留样式、页眉、页脚。"""
        body = doc.element.body
        to_remove = [child for child in body if child.tag != qn("w:sectPr")]
        for child in to_remove:
            body.remove(child)

    @staticmethod
    def _set_margins(doc: Any) -> None:
        for section in doc.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(2.8)
            section.right_margin = Cm(2.8)

    @staticmethod
    def _add_header_paragraphs(doc: Any, display_date: str) -> None:
        """添加 航天通报（首页表头_中）和日期（首页表头_下）。

        若模板中存在这两个样式则直接应用，否则以 Normal 居中替代。
        """
        style_names = {s.name for s in doc.styles}

        # 航天通报
        if "首页表头_中" in style_names:
            p_title = doc.add_paragraph(style="首页表头_中")
        else:
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.add_run("航天通报")

        # 日期
        if "首页表头_下" in style_names:
            p_date = doc.add_paragraph(style="首页表头_下")
        else:
            p_date = doc.add_paragraph()
            p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_date.add_run(display_date)

    @staticmethod
    def _add_body_text(doc: Any, body: str) -> None:
        """添加正文段落（16pt，首行缩进2字符）。"""
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(32)  # 2 × 16pt
        run = p.add_run(body)
        run.font.size = Pt(16)

    @staticmethod
    def _add_attachment_index(doc: Any, region_name: str, include_image: bool = True) -> None:
        """在正文末尾添加附件引用列表（Normal 样式）。"""
        doc.add_paragraph()  # 空行
        attachments = []
        if include_image:
            attachments.append((1, "目标分布图"))
        # 附件编号随 include_image 动态偏移
        base = 2 if include_image else 1
        attachments.append((base,     "组成分布统计表"))
        attachments.append((base + 1, "装备分布统计表"))
        for idx, suffix in attachments:
            p = doc.add_paragraph()
            p.add_run(f"附件{idx}：{region_name}{suffix}")

    @staticmethod
    def _add_attachment1_image(
        doc: Any, image_path: str | None, region_name: str
    ) -> None:
        """附件1：侦察图（独立页，全页宽）。"""
        doc.add_page_break()
        p_head = doc.add_paragraph()
        p_head.add_run("附件1：")

        if image_path and Path(image_path).exists():
            try:
                doc.add_picture(image_path, width=Cm(15.4))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                doc.add_paragraph().add_run(f"[图像加载失败：{image_path}]").italic = True
        else:
            doc.add_paragraph().add_run("（侦察图像暂无）").italic = True

        cap = doc.add_paragraph(style="Caption") if "Caption" in {
            s.name for s in doc.styles
        } else doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(f"图 1  {region_name}目标分布图")

    @staticmethod
    def _add_attachment2_table(
        doc: Any, rows: list[dict], region_name: str
    ) -> None:
        """附件2：组成分布统计表（4列）。"""
        doc.add_paragraph()
        p_head = doc.add_paragraph()
        p_head.add_run("附件2：")

        DocxAssembler._add_distribution_table(
            doc=doc,
            rows=rows,
            region_name=region_name,
            table_no=1,
            suffix="组成分布统计表",
            empty_text="（暂无目标数据）",
            type_getter=lambda row_data: (
                row_data.get("sub_type")
                or row_data.get("target_type")
                or row_data.get("equipment_type", "未知")
            ),
        )
    @staticmethod
    def _add_attachment3_table(
        doc: Any, rows: list[dict], region_name: str
    ) -> None:
        """附件3：装备分布统计表（4列）。"""
        doc.add_paragraph()
        p_head = doc.add_paragraph()
        p_head.add_run("附件3：")

        DocxAssembler._add_distribution_table(
            doc=doc,
            rows=rows,
            region_name=region_name,
            table_no=1,
            suffix="装备分布统计表",
            empty_text="（暂无装备数据）",
            type_getter=lambda row_data: (
                row_data.get("equipment_type")
                or row_data.get("sub_type")
                or row_data.get("target_type", "未知")
            ),
        )

    @staticmethod
    def _add_distribution_table(
        doc: Any,
        rows: list[dict],
        region_name: str,
        table_no: int,
        suffix: str,
        empty_text: str,
        type_getter: Any,
    ) -> None:
        if not rows:
            doc.add_paragraph(empty_text).italic = True
            return

        cap_style = "Caption" if "Caption" in {s.name for s in doc.styles} else "Normal"
        chunks = _chunk_table_rows(rows)
        for chunk_index, chunk in enumerate(chunks):
            if chunk_index:
                doc.add_page_break()

            cap = doc.add_paragraph(style=cap_style)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(6)
            continuation = f"（续{chunk_index}）" if chunk_index else ""
            cap.add_run(f"表 {table_no}{continuation}  {region_name}{suffix}")

            table = doc.add_table(rows=1 + len(chunk), cols=4)
            table.style = "Table Grid"

            hdr = table.rows[0].cells
            for i, h in enumerate(_TABLE_HEADERS):
                hdr[i].text = h
                run = hdr[i].paragraphs[0].runs
                r = run[0] if run else hdr[i].paragraphs[0].add_run(h)
                r.bold = True
                r.font.size = Pt(11)

            for ri, row_data in enumerate(chunk, start=1):
                cells = table.rows[ri].cells
                cells[0].text = str(row_data.get("seq", ri))
                cells[1].text = str(type_getter(row_data))
                lon = row_data.get("lon")
                lat = row_data.get("lat")
                cells[2].text = _format_coordinate(lon, "lon")
                cells[3].text = _format_coordinate(lat, "lat")
                for cell in cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(10)

    @staticmethod
    def _add_footer(
        doc: Any,
        satellite: str,
        sensor: str,
        acq_time: str,
        pipeline_run_id: str,
        acceptance_run_id: str,
        draft_mode: bool,
    ) -> None:
        for section in doc.sections:
            footer = section.footer
            for para in footer.paragraphs:
                for run in para.runs:
                    run.text = ""
            if not footer.paragraphs:
                footer.add_paragraph()

            p = footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parts: list[str] = []
            if satellite:
                parts.append(f"卫星：{satellite}")
            if sensor:
                parts.append(f"传感器：{sensor}")
            if acq_time:
                # Format to YYYY-MM-DD for footer
                date_part = acq_time[:10] if len(acq_time) >= 10 else acq_time
                parts.append(f"成像时间：{date_part}")
            if pipeline_run_id:
                parts.append(f"流水线：{pipeline_run_id}")
            if acceptance_run_id:
                parts.append(f"验收批次：{acceptance_run_id}")

            footer_text = " | ".join(parts)
            if draft_mode:
                footer_text += "  【自动生成草稿，待人工审核】"

            run = p.add_run(footer_text)
            run.font.size = Pt(9)
            if draft_mode:
                run.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)

    @staticmethod
    def _strip_unreferenced_document_images(docx_path: str) -> None:
        """Remove stale template image rels that are no longer referenced."""
        rels_name = "word/_rels/document.xml.rels"
        document_name = "word/document.xml"
        path = Path(docx_path)
        tmp_path = path.with_name(path.stem + ".tmp" + path.suffix)

        try:
            with ZipFile(path, "r") as zin:
                if rels_name not in zin.namelist() or document_name not in zin.namelist():
                    return
                document_xml = zin.read(document_name).decode("utf-8", errors="ignore")
                used_rel_ids = set(re.findall(r'r:(?:embed|link)="([^"]+)"', document_xml))

                rels_root = ET.fromstring(zin.read(rels_name))
                rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
                removed_targets: set[str] = set()
                for rel in list(rels_root):
                    rel_id = rel.attrib.get("Id")
                    rel_type = rel.attrib.get("Type", "")
                    target = rel.attrib.get("Target", "")
                    if (
                        rel_id
                        and rel_id not in used_rel_ids
                        and rel_type.endswith("/image")
                        and target.startswith("media/")
                    ):
                        rels_root.remove(rel)
                        removed_targets.add("word/" + target)

                if not removed_targets:
                    return

                rels_xml = ET.tostring(
                    rels_root,
                    encoding="utf-8",
                    xml_declaration=True,
                    short_empty_elements=True,
                )
                with ZipFile(tmp_path, "w", compression=ZIP_DEFLATED) as zout:
                    for item in zin.infolist():
                        if item.filename in removed_targets:
                            continue
                        data = rels_xml if item.filename == rels_name else zin.read(item.filename)
                        zout.writestr(item, data)
            tmp_path.replace(path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()

    @staticmethod
    def patch_existing_docx(
        docx_path: str,
        evidence_package: dict,
        draft_mode: bool = True,
    ) -> str:
        """Patch an existing docx when python-docx is unavailable.

        This updates the main body paragraph and footer text in-place via XML.
        """
        path = Path(docx_path)
        if not path.exists():
            raise FileNotFoundError(f"docx file not found for patching: {docx_path}")

        report = evidence_package.get("report", {})
        body_value = report.get("body")
        if not isinstance(body_value, str) or not body_value.strip():
            raise ValueError("report.body is required before patching an existing DOCX.")
        body_text = body_value
        inp = evidence_package.get("input", {})
        metadata = inp.get("metadata", {})
        trace = evidence_package.get("trace", {})
        footer_text = DocxAssembler._build_footer_text(
            satellite=str(metadata.get("satellite", "")),
            sensor=str(metadata.get("sensor", "")),
            acq_time=str(metadata.get("acquisition_time", "")),
            pipeline_run_id=str(trace.get("pipeline_run_id", "")),
            acceptance_run_id=str(trace.get("acceptance_run_id", "")),
            draft_mode=draft_mode,
        )

        tmp_path = path.with_name(path.stem + ".patch" + path.suffix)
        try:
            with ZipFile(path, "r") as zin, ZipFile(tmp_path, "w", compression=ZIP_DEFLATED) as zout:
                footer_names = [
                    name for name in zin.namelist()
                    if name.startswith("word/footer") and name.endswith(".xml")
                ]
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "word/document.xml":
                        data = DocxAssembler._patch_document_xml(data, body_text)
                    elif item.filename in footer_names:
                        data = DocxAssembler._patch_footer_xml(data, footer_text)
                    zout.writestr(item, data)
            tmp_path.replace(path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()
        return str(path)

    # ------------------------------------------------------------------
    # 辅助方法
    # ------------------------------------------------------------------

    @staticmethod
    def _format_date(report_date: str, acquisition_time: str) -> str:
        """将 ISO 日期或报告日期格式化为 YYYY年MM月DD日。"""
        for src in (report_date, acquisition_time):
            if not src:
                continue
            # Already in Chinese format
            if "年" in src:
                return src
            # ISO date prefix YYYY-MM-DD
            m = re.match(r"(\d{4})-(\d{2})-(\d{2})", src)
            if m:
                return f"{m.group(1)}年{int(m.group(2))}月{int(m.group(3))}日"
        return datetime.now(timezone.utc).strftime("%Y年%-m月%-d日")

    @staticmethod
    def _build_footer_text(
        satellite: str,
        sensor: str,
        acq_time: str,
        pipeline_run_id: str,
        acceptance_run_id: str,
        draft_mode: bool,
    ) -> str:
        parts: list[str] = []
        if satellite:
            parts.append(f"卫星：{satellite}")
        if sensor:
            parts.append(f"传感器：{sensor}")
        if acq_time:
            date_part = acq_time[:10] if len(acq_time) >= 10 else acq_time
            parts.append(f"成像时间：{date_part}")
        if pipeline_run_id:
            parts.append(f"流水线：{pipeline_run_id}")
        if acceptance_run_id:
            parts.append(f"验收批次：{acceptance_run_id}")

        footer_text = " | ".join(parts)
        if draft_mode:
            footer_text += "  【自动生成草稿，待人工审核】"
        return footer_text

    @staticmethod
    def _patch_document_xml(data: bytes, body_text: str) -> bytes:
        root = ET.fromstring(data)
        body = root.find("w:body", _XML_NS)
        if body is None:
            return data

        seen_title = False
        seen_date = False
        for para in body.findall("w:p", _XML_NS):
            text = "".join(node.text or "" for node in para.findall(".//w:t", _XML_NS)).strip()
            if not text:
                continue
            if not seen_title and text == "航天通报":
                seen_title = True
                continue
            if seen_title and not seen_date and re.match(r"\d{4}年\d{1,2}月\d{1,2}日", text):
                seen_date = True
                continue
            if text.startswith("附件"):
                break
            DocxAssembler._replace_paragraph_text(para, body_text)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return data

    @staticmethod
    def _patch_footer_xml(data: bytes, footer_text: str) -> bytes:
        root = ET.fromstring(data)
        para = root.find("w:p", _XML_NS)
        if para is None:
            para = ET.SubElement(root, f"{{{_WORD_NS}}}p")
        DocxAssembler._replace_paragraph_text(
            para,
            footer_text,
            font_half_points=18,
            color="AA0000",
        )
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _replace_paragraph_text(
        para: ET.Element,
        text: str,
        font_half_points: int | None = None,
        color: str | None = None,
    ) -> None:
        ppr = para.find("w:pPr", _XML_NS)
        for child in list(para):
            if child is not ppr:
                para.remove(child)

        run = ET.SubElement(para, f"{{{_WORD_NS}}}r")
        if font_half_points is not None or color is not None:
            rpr = ET.SubElement(run, f"{{{_WORD_NS}}}rPr")
            if font_half_points is not None:
                ET.SubElement(rpr, f"{{{_WORD_NS}}}sz", {f"{{{_WORD_NS}}}val": str(font_half_points)})
                ET.SubElement(rpr, f"{{{_WORD_NS}}}szCs", {f"{{{_WORD_NS}}}val": str(font_half_points)})
            if color is not None:
                ET.SubElement(rpr, f"{{{_WORD_NS}}}color", {f"{{{_WORD_NS}}}val": color})
        t = ET.SubElement(run, f"{{{_WORD_NS}}}t")
        t.text = text


# ---------------------------------------------------------------------------
# 顶层便捷函数
# ---------------------------------------------------------------------------

def assemble_docx(
    evidence_package: dict,
    output_dir: str,
    template_path: str | None = None,
    draft_mode: bool = True,
    include_image: bool = True,
) -> str:
    return DocxAssembler().assemble(
        evidence_package,
        output_dir,
        template_path=template_path,
        draft_mode=draft_mode,
        include_image=include_image,
    )
